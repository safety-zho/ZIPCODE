"""엑셀 읽기·쓰기. 원본 서식을 유지하려고 openpyxl로 원본 통합문서에 바로 값을 쓴다."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_SCAN_ROWS = 30
ADDRESS_EXCLUDE = ("상세", "이메일", "메일", "email", "e-mail", "ip", "url", "홈페이지", "영문")


class ExcelError(Exception):
    pass


@dataclass
class Header:
    col: int
    letter: str
    title: str


@dataclass
class Layout:
    sheets: list[str]
    sheet: str
    header_row: int
    headers: list[Header]
    address_col: int | None
    zip_col: int | None  # None이면 새 열을 만든다
    new_zip_col: int  # 우편번호 열이 없을 때 만들 위치
    rows: list[dict] = field(default_factory=list)
    truncated: bool = False

    def to_dict(self) -> dict:
        return {
            "sheets": self.sheets,
            "sheet": self.sheet,
            "header_row": self.header_row,
            "headers": [h.__dict__ for h in self.headers],
            "address_col": self.address_col,
            "zip_col": self.zip_col,
            "new_zip_col": self.new_zip_col,
            "new_zip_letter": get_column_letter(self.new_zip_col),
            "rows": self.rows,
            "truncated": self.truncated,
        }


def _txt(v) -> str:
    if v is None:
        return ""
    return " ".join(str(v).replace(" ", " ").split())


def _key(v) -> str:
    return re.sub(r"\s+", "", _txt(v)).lower()


def open_workbook(data: bytes, data_only: bool):
    try:
        return load_workbook(io.BytesIO(data), data_only=data_only)
    except Exception as e:  # zipfile.BadZipFile, InvalidFileException 등
        raise ExcelError(
            "엑셀 파일을 열 수 없습니다. .xlsx 형식인지 확인해 주세요. "
            "(.xls 파일은 엑셀에서 '다른 이름으로 저장 → Excel 통합 문서(.xlsx)'로 바꾼 뒤 올려 주세요.)"
        ) from e


def _address_score(title: str) -> int:
    k = _key(title)
    if "주소" not in k:
        return 0
    if any(x in k for x in ADDRESS_EXCLUDE):
        return 0
    if k in ("주소", "도로명주소", "주소(도로명)", "도로명", "소재지", "주소지"):
        return 3
    if k.startswith("도로명") or k.endswith("주소"):
        return 2
    return 1


def _is_zip_title(title: str) -> bool:
    k = _key(title)
    return "우편번호" in k or k in ("우편", "zip", "zipcode", "postcode", "우편번호(5자리)")


def _find_header_row(ws: Worksheet) -> int | None:
    best = None
    for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
        for c in range(1, ws.max_column + 1):
            score = _address_score(ws.cell(r, c).value)
            if score and (best is None or score > best[0]):
                best = (score, r)
        if best and best[0] == 3:
            break
    return best[1] if best else None


def _first_nonempty_row(ws: Worksheet) -> int:
    for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
        if any(_txt(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)):
            return r
    return 1


def inspect(
    data: bytes,
    sheet: str | None = None,
    address_col: int | None = None,
    zip_col: int | None = None,
    use_new_zip_col: bool = False,
    max_rows: int = 5000,
) -> Layout:
    wb = open_workbook(data, data_only=True)
    sheets = wb.sheetnames

    if sheet and sheet in sheets:
        ws = wb[sheet]
    else:
        # '주소' 제목이 있는 첫 시트를 고른다
        ws = next((wb[n] for n in sheets if _find_header_row(wb[n]) is not None), wb[sheets[0]])

    header_row = _find_header_row(ws) or _first_nonempty_row(ws)
    headers: list[Header] = []
    last_col = 0
    for c in range(1, ws.max_column + 1):
        t = _txt(ws.cell(header_row, c).value)
        if t:
            headers.append(Header(c, get_column_letter(c), t))
            last_col = c
    if not headers:
        raise ExcelError("시트에서 열 제목을 찾지 못했습니다. 첫 행에 '주소', '우편번호' 같은 제목이 있는지 확인해 주세요.")

    if address_col is None:
        scored = [(h, _address_score(h.title)) for h in headers]
        scored = [s for s in scored if s[1]]
        if scored:
            address_col = max(scored, key=lambda s: s[1])[0].col

    if use_new_zip_col:
        zip_col = None
    elif zip_col is None:
        z = next((h for h in headers if _is_zip_title(h.title) and h.col != address_col), None)
        zip_col = z.col if z else None

    if address_col is not None and zip_col == address_col:
        raise ExcelError("주소 열과 우편번호 열이 같습니다. 서로 다른 열을 골라 주세요.")

    layout = Layout(
        sheets=sheets,
        sheet=ws.title,
        header_row=header_row,
        headers=headers,
        address_col=address_col,
        zip_col=zip_col,
        new_zip_col=max(last_col, ws.max_column) + 1,
    )

    if address_col is None:
        return layout

    for r in range(header_row + 1, ws.max_row + 1):
        addr = _txt(ws.cell(r, address_col).value)
        if not addr:
            continue
        if len(layout.rows) >= max_rows:
            layout.truncated = True
            break
        existing = _txt(ws.cell(r, zip_col).value) if zip_col else ""
        layout.rows.append({"row": r, "address": addr, "existing": existing})
    return layout


def fill(
    data: bytes,
    sheet: str,
    address_col: int,
    zip_col: int | None,
    assignments: dict[int, dict],
    clear_unresolved: bool,
    add_check_column: bool,
    max_rows: int = 5000,
) -> tuple[bytes, dict]:
    """assignments: {행번호: {"zip": "06232", "road": "서울특별시 ..."}} — 서명 검증을 마친 값만 들어온다."""
    layout = inspect(data, sheet=sheet, address_col=address_col, zip_col=zip_col,
                     use_new_zip_col=zip_col is None, max_rows=max_rows)
    valid_rows = {r["row"] for r in layout.rows}
    bad = [r for r in assignments if r not in valid_rows]
    if bad:
        raise ExcelError(f"주소가 없는 행에는 우편번호를 쓸 수 없습니다: {', '.join(map(str, bad[:5]))}행")

    wb = open_workbook(data, data_only=False)
    ws = wb[layout.sheet]
    hr = layout.header_row

    if zip_col is None:
        zip_col = layout.new_zip_col
        _copy_header_style(ws, hr, address_col, zip_col)
        ws.cell(hr, zip_col).value = "우편번호"
        ws.column_dimensions[get_column_letter(zip_col)].width = 10

    check_col = None
    if add_check_column:
        check_col = max(ws.max_column, zip_col) + 1
        _copy_header_style(ws, hr, address_col, check_col)
        ws.cell(hr, check_col).value = "인터넷우체국 검색 주소"
        ws.column_dimensions[get_column_letter(check_col)].width = 48

    filled = cleared = 0
    for r in sorted(valid_rows):
        cell = ws.cell(r, zip_col)
        a = assignments.get(r)
        if a:
            cell.value = a["zip"]
            cell.number_format = "@"  # 앞자리 0 유지 (06232)
            filled += 1
            if check_col:
                ws.cell(r, check_col).value = a.get("road", "")
        elif clear_unresolved and cell.value not in (None, ""):
            cell.value = None
            cleared += 1

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), {
        "filled": filled,
        "cleared": cleared,
        "unresolved": len(valid_rows) - filled,
        "zip_letter": get_column_letter(zip_col),
    }


def _copy_header_style(ws: Worksheet, row: int, src_col: int, dst_col: int) -> None:
    from copy import copy

    src = ws.cell(row, src_col)
    dst = ws.cell(row, dst_col)
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
